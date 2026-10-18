# -*- coding: utf-8 -*-
"""
Копия работы студента с отметками: абзацы с замечаниями подсвечены, а на полях —
примечания Word («Рецензирование → Примечания») с текстом замечаний. Текст работы
не меняется. Одно примечание на абзац: в нём все замечания к этому абзацу.

Место каждого замечания берётся из Issue.anchor (см. report.py), которое проверка
заполняет порядковым номером блока в iter_blocks() — при повторном открытии файла
порядок тот же, поэтому отметки попадают ровно туда, где найдена ошибка.
"""

from collections import OrderedDict

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.run import Run

from .docx_utils import iter_blocks

AUTHOR = "Нормоконтроль"
INITIALS = "НК"
HIGHLIGHT = WD_COLOR_INDEX.YELLOW


def _anchor_runs(p):
    """
    Run'ы абзаца, к которым можно привязать примечание. Берём только прямые потомки
    абзаца (не внутри гиперссылок/полей) — так метки начала и конца примечания
    гарантированно стоят на правильном уровне. В пустом абзаце (например, в лишней
    пустой строке) добавляем пустой run: текст при этом не меняется.
    """
    runs = [r for r in p.runs if r.text.strip()] or list(p.runs)
    if not runs:
        runs = [p.add_run("")]
    return runs


def _highlight(p):
    for r in p._p.iter(qn("w:r")):
        if r.find(qn("w:t")) is not None:   # рисунки и поля без текста не трогаем
            Run(r, p).font.highlight_color = HIGHLIGHT


def _comment(doc, runs, issues, title=None):
    lines = ([title] if title else []) + [f"• {i.category}: {i.message}" for i in issues]
    c = doc.add_comment(runs, text=lines[0], author=AUTHOR, initials=INITIALS)
    for line in lines[1:]:
        c.add_paragraph(line)
    return c


def _section_starts(blocks):
    """Номер блока, с которого начинается каждый раздел документа (по w:sectPr в абзацах)."""
    starts = [0]
    for i, (kind, obj) in enumerate(blocks):
        if kind == "p" and obj._p.pPr is not None and obj._p.pPr.find(qn("w:sectPr")) is not None:
            starts.append(i + 1)
    return starts


def _first_text_paragraph(blocks, start=0):
    for kind, obj in blocks[start:]:
        if kind == "p" and obj.text.strip():
            return obj
    for kind, obj in blocks[start:]:
        if kind == "p":
            return obj
    return None


def annotate_docx(src_path, dst_path, report):
    """
    Сохраняет в dst_path копию src_path с отметками замечаний из report.
    Возвращает число добавленных примечаний.
    """
    doc = Document(src_path)
    blocks = list(iter_blocks(doc))
    section_starts = _section_starts(blocks)

    # группируем замечания по месту, сохраняя порядок появления
    by_place = OrderedDict()
    general = []
    for issue in report.issues:
        a = issue.anchor
        if a and a[0] == "p" and a[1] < len(blocks) and blocks[a[1]][0] == "p":
            by_place.setdefault(("p", a[1]), []).append(issue)
        elif a and a[0] == "cell" and a[1] < len(blocks) and blocks[a[1]][0] == "tbl":
            by_place.setdefault(a, []).append(issue)
        elif a and a[0] == "section" and a[1] < len(section_starts):
            by_place.setdefault(("section", a[1]), []).append(issue)
        else:
            general.append(issue)

    count = 0
    for place, issues in by_place.items():
        kind = place[0]
        if kind == "p":
            p = blocks[place[1]][1]
            title = None
        elif kind == "cell":
            _, bi, ri, ci = place
            try:
                cell = blocks[bi][1].rows[ri].cells[ci]
            except IndexError:
                general.extend(issues)
                continue
            p = cell.paragraphs[0]
            title = None
        else:  # section
            p = _first_text_paragraph(blocks, section_starts[place[1]])
            if p is None:
                general.extend(issues)
                continue
            title = f"Раздел документа №{place[1] + 1}:"
        _highlight(p)
        _comment(doc, _anchor_runs(p), issues, title)
        count += 1

    if general:
        p = _first_text_paragraph(blocks)
        if p is not None:
            _comment(doc, _anchor_runs(p), general, "Замечания ко всему документу:")
            count += 1

    doc.save(dst_path)
    return count
