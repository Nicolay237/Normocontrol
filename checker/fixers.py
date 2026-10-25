# -*- coding: utf-8 -*-
"""
Автоисправление .docx по профилю норм.

Принципы:
  * Меняется только оформление (свойства абзацев и run'ов), смысл текста не
    трогается. Исключения — перечислены в отчёте «что изменено»: регистр
    ПРОПИСНЫХ заголовков, точка в конце заголовка/подписи, пробелы,
    маркер списка → дефис, формат «Рисунок N – Название».
  * Исправляется только то, что нарушает норму — остальное остаётся как есть.
  * Правила те же, что у проверки (DocxChecker): после исправления файл
    проверяется заново, и всё, что починить нельзя, остаётся в отчёте.
  * Исходный файл не изменяется — результат пишется в новый.
"""

import re
from collections import Counter
from dataclasses import dataclass, field

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.text.paragraph import Paragraph

from .docx_checks import (
    DocxChecker, ALT_FONTS, HEADING_NUMBER_PREFIX_RE, BULLET_MARKERS, CHAPTER_NUMBER_RE, FIG_RE, TAB_RE,
    DASHES, STRUCT_APPENDIX_RE, _toc_level, ends_with_sentence_dot,
)
from .docx_utils import get_all_runs, run_in_hyperlink, has_image, has_page_break, emu_to_cm
from .report import Report
from . import bibliography as bib
from .typography import SPACE_BEFORE_SIMPLE, SPACE_BEFORE_DOT, MULTI_SPACE, HYPHEN_AS_DASH, CODE_LIKE

TNR = "Times New Roman"

# Категории замечаний, которые исправитель устраняет сам (для оценки «исправим около N»).
# Не исправляются: пропуски в нумерации рисунков/таблиц и отсутствующие разделы.
FIXABLE_CATEGORIES = {
    "Формат листа", "Поля страницы", "Нумерация страниц", "Шрифт", "Шрифт (таблица)", "Кегль шрифта",
    "Цвет шрифта", "Межстрочный интервал", "Выравнивание", "Красная строка",
    "Подписи к рисункам и таблицам", "Оформление заголовка", "Оформление оглавления",
    "Перечисления", "Пробелы", "Пустые строки",
}


# Замечания к списку источников, которые исправитель правит сам (начало текста замечания)
FIXABLE_BIB_MESSAGES = (
    "после фамилии автора ставится запятая", "инициалы пишутся через пробел", "место издания пишется полностью",
    "место издания и издательство разделяются", "вместо дефиса с пробелами", "после знака «№»",
    "перед «URL:» ставятся точка и тире",
)
FIXABLE_CODE_MESSAGES = ("листинг выравнивается по левому краю", "у строк листинга не должно быть абзацного отступа")


def estimate_fixable(report):
    """Сколько замечаний отчёта исправитель, скорее всего, устранит (оценка до запуска)."""
    return sum(1 for i in report.issues
               if i.category in FIXABLE_CATEGORIES
               or (i.category == "Список источников" and i.message.startswith(FIXABLE_BIB_MESSAGES))
               or (i.category == "Листинги" and i.message.startswith(FIXABLE_CODE_MESSAGES)))


@dataclass
class FixResult:
    """
    Итог исправления. Пользователю показываются ЗАМЕЧАНИЯ (как в отчёте проверки),
    а не отдельные правки XML: одно замечание может потребовать нескольких правок
    (шрифт в каждом run'е абзаца), поэтому число правок с числом замечаний не совпадает.
    """
    applied: dict = field(default_factory=dict)   # {«что сделано»: сколько правок} — для журнала
    before: Report = None
    after: Report = None

    def by_category(self):
        """[{name, before, after, fixed}] в порядке категорий исходного отчёта."""
        b = {k: len(v) for k, v in self.before.by_category().items()}
        a = {k: len(v) for k, v in self.after.by_category().items()}
        rows = [{"name": k, "before": n, "after": a.get(k, 0), "fixed": max(0, n - a.get(k, 0))} for k, n in b.items()]
        rows += [{"name": k, "before": 0, "after": n, "fixed": 0} for k, n in a.items() if k not in b]
        return rows

    @property
    def total_fixed(self):
        return sum(r["fixed"] for r in self.by_category())

    @property
    def total_new(self):
        """Замечания, которых не было до исправления (по категориям)."""
        return sum(max(0, r["after"] - r["before"]) for r in self.by_category())

    def to_dict(self):
        return {
            "issues_before": len(self.before.issues),
            "issues_after": len(self.after.issues),
            "total_fixed": self.total_fixed,
            "total_new": self.total_new,
            "categories": self.by_category(),
            "operations": [{"name": k, "count": v} for k, v in self.applied.items()],
            "remaining": self.after.to_dict()["categories"],
            "notes": self.after.notes,
        }


# ── низкоуровневые правки XML ─────────────────────────────────────────────

def _set_font_name(run, name=TNR):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in list(rf.attrib):  # тема шрифта перекрывает явное имя — убираем
        del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def _set_font_size(run, size_pt):
    run.font.size = Pt(size_pt)
    rpr = run._r.get_or_add_rPr()
    cs = rpr.find(qn("w:szCs"))
    if cs is None:
        cs = OxmlElement("w:szCs")
        rpr.append(cs)
    cs.set(qn("w:val"), str(int(round(size_pt * 2))))


def _set_black(run):
    run.font.color.rgb = RGBColor(0, 0, 0)
    c = run._r.get_or_add_rPr().find(qn("w:color"))
    for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
        if c is not None and c.get(qn(a)) is not None:
            del c.attrib[qn(a)]


def _set_first_line(p, cm):
    p.paragraph_format.first_line_indent = Cm(cm)
    ind = p._p.pPr.find(qn("w:ind"))
    for a in ("w:firstLineChars", "w:hangingChars"):  # Word отдаёт им приоритет над firstLine
        if ind is not None and ind.get(qn(a)) is not None:
            del ind.attrib[qn(a)]


def _wts(p):
    return [t for t in p._p.iter(qn("w:t"))]


def _replace_spans(ts, edits):
    """
    Применяет правки (начало, конец, текст) к склеенному тексту элементов w:t,
    сохраняя разбиение на run'ы: новый текст попадает в run, где начинается правка,
    удалённые символы вырезаются из всех задетых run'ов. Пересекающиеся правки отбрасываются.
    """
    if not edits:
        return False
    starts, pos = [], 0
    for t in ts:
        starts.append(pos)
        pos += len(t.text or "")
    chosen, last_start = [], None
    for a, b, new in sorted(edits, key=lambda e: (e[0], e[1]), reverse=True):
        if last_start is not None and b > last_start:
            continue
        chosen.append((a, b, new))
        last_start = a
    for a, b, new in chosen:          # с конца, чтобы смещения более ранних правок не менялись
        idx = next((i for i in range(len(ts) - 1, -1, -1) if starts[i] <= a), 0)
        for i in range(len(ts)):
            text = ts[i].text or ""
            lo, hi = starts[i], starts[i] + len(text)
            cut_a, cut_b = max(a, lo), min(b, hi)
            piece_new = new if i == idx else ""
            if i == idx or (cut_a < cut_b):
                ka, kb = cut_a - lo, max(cut_a, cut_b) - lo
                if i == idx and not (lo <= a <= hi):
                    continue
                ts[i].text = text[:ka] + piece_new + text[kb:]
                if piece_new:
                    ts[i].set(qn("xml:space"), "preserve")
    return True


def _has_field(p):
    return bool(p._p.xpath(".//w:fldChar | .//w:fldSimple | .//w:instrText"))


def _add_page_field(paragraph):
    run = paragraph.add_run()
    for kind in ("begin", "instr", "separate", "text", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def _normalize_caption(text, full):
    """«Рис. 2. схема.» → «Рисунок 2 – Схема» (None, если формат не распознан)."""
    rx = FIG_RE if full == "Рисунок" else TAB_RE
    m = rx.match(text.strip())
    if not m:
        return None
    _, _, num, rest = m.groups()
    new = f"{full} {num}"
    name = rest.strip().lstrip(DASHES + ".:").strip()
    if name:
        if ends_with_sentence_dot(name):
            name = name.rstrip()[:-1].rstrip()
        name = name[:1].upper() + name[1:]
        new += f" – {name}"
    return new


# ── исправитель ───────────────────────────────────────────────────────────

class _Fixer:
    def __init__(self, doc, norms):
        self.doc, self.n = doc, norms
        self.chk = DocxChecker(doc, Report(), norms)
        self.res = self.chk.res
        self.applied = Counter()

    def done(self, label, k=1):
        self.applied[label] += k

    # ---- запуск
    def run(self):
        items = self.chk.prepare()
        self._bib_ids = {id(p._p) for p in bib.find_section(self.chk, items)[1]}
        self.fix_sections()
        self.fix_page_numbers()
        prev_formula = False
        for it in items:
            if it["kind"] == "tbl":
                if it["role"] != "front":   # таблица подписей на титульном листе — по своей форме
                    self.fix_table(it["obj"])
                prev_formula = False
                continue
            p, role = it["obj"], it["role"]
            if role == "formula":
                prev_formula = True
                continue
            if role in ("body", "list", "figure_caption", "table_caption"):
                self.fix_runs(p, floor_only=role.endswith("caption"))
                self.fix_body_paragraph(p, role, prev_formula)
            elif role in ("heading", "structural"):
                self.fix_runs(p, heading=True)
                self.fix_heading(p, role)
            elif role == "toc":
                self.fix_toc(p)
            elif role == "code":
                # листинг: шрифт, отступы и интервалы кода не трогаем — только выравнивание по левому краю
                if self.res.alignment(p) in ("justify", "center"):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    self.done("Листинг выровнен по левому краю")
                if self.res.first_line_indent_cm(p, 12) > self.n["first_line_indent_tolerance_cm"]:
                    _set_first_line(p, 0)
                    self.done("Убран абзацный отступ в листинге")
            elif role == "appendix_title":
                # заголовок приложения: по центру, без точки в конце (ГОСТ 7.32-2017, 6.17)
                if self.res.alignment(p) != "center":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _set_first_line(p, 0)
                    self.done("Заголовок приложения выровнен по центру")
                if ends_with_sentence_dot(p.text.strip()):
                    for t in reversed([t for t in _wts(p) if (t.text or "").strip()]):
                        t.text = t.text.rstrip()[:-1]
                        break
                    self.done("Убрана точка в конце заголовка")
            if role not in ("empty", "image"):
                prev_formula = False
            if role in ("body", "list", "heading", "structural", "figure_caption", "table_caption"):
                self.fix_text(p, role)
        for it in list(self.chk.headings_missing_new_page(items)):
            it["obj"].paragraph_format.page_break_before = True
            self.done("Заголовок начат с новой страницы")
        if self.move_misplaced_captions(items):
            items = self.chk.prepare()   # абзацы переехали — пустые строки ищем по новому порядку
        self.remove_extra_empty(items)
        if self.n["bib_standard"] == "7.0.100":   # исправитель знает правила только ГОСТ Р 7.0.100
            self.fix_bibliography(items)

    # ---- страница
    def fix_sections(self):
        n, tol = self.n, self.n["margins_tolerance_cm"]
        w0, h0 = n["page_size_cm"]
        for s in self.doc.sections:
            w, h = emu_to_cm(s.page_width), emu_to_cm(s.page_height)
            if not (w and h):
                continue
            landscape = w > h
            a, b = sorted((w, h))
            ok_a4 = abs(a - w0) <= n["page_size_tolerance_cm"] and abs(b - h0) <= n["page_size_tolerance_cm"]
            ok_a3 = abs(a - 29.7) <= 0.5 and abs(b - 42.0) <= 0.5
            if not (ok_a4 or ok_a3):
                pw, ph = (Cm(h0), Cm(w0)) if landscape else (Cm(w0), Cm(h0))
                s.page_width, s.page_height = pw, ph
                self.done("Формат листа приведён к A4")
            if landscape:
                continue
            for side, (lo, hi) in n["margins_cm"].items():
                attr = f"{side}_margin"
                cur = emu_to_cm(getattr(s, attr))
                if cur is None:
                    continue
                if cur < lo - tol or cur > hi + tol:
                    setattr(s, attr, Cm(lo if cur < lo else hi))
                    self.done("Поля страницы исправлены")

    def fix_page_numbers(self):
        if not self.n["check_page_numbers"]:
            return
        for i, s in enumerate(self.doc.sections):
            footer = s.footer
            if i > 0 and footer.is_linked_to_previous:
                continue
            page_paras = [Paragraph(el, footer) for el in footer._element.iter(qn("w:p"))
                          if "PAGE" in el.xml and ("instrText" in el.xml or "fldSimple" in el.xml)]
            if not page_paras:
                if any(self._has_page_in_header(sec) for sec in self.doc.sections):
                    continue  # номер в верхнем колонтитуле — переносить автоматически не беремся
                p = footer.paragraphs[0] if footer.paragraphs and not footer.paragraphs[0].text.strip() \
                    else footer.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_page_field(p)
                self.done("Добавлена нумерация страниц (внизу по центру)")
            else:
                for p in page_paras:
                    if self.res.alignment(p) != "center" and "<w:tab/>" not in p._p.xml:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        self.done("Номер страницы выровнен по центру")

    @staticmethod
    def _has_page_in_header(sec):
        h = sec.header
        return any("PAGE" in el.xml and "instrText" in el.xml for el in h._element.iter(qn("w:p")))

    # ---- run'ы
    def fix_runs(self, p, floor_only=False, heading=False):
        n, res = self.n, self.res
        tol = n["font_size_tolerance_pt"]
        for run in self.chk._runs(p):
            if res.script(run, p):
                continue
            fname = res.font_name(run, p)
            if fname and fname != n["font_name"] and fname.lower() not in ALT_FONTS:
                _set_font_name(run, n["font_name"])
                self.done("Шрифт заменён на " + n["font_name"])
            size = res.font_size(run, p)
            if size is not None:
                if floor_only:
                    if size < n["font_size_floor_pt"] - tol:
                        _set_font_size(run, n["font_size_floor_pt"])
                        self.done("Кегль приведён к норме")
                elif size < n["font_size_min_pt"] - tol or size > n["font_size_max_pt"] + tol:
                    _set_font_size(run, n["font_size_pt"])
                    self.done("Кегль приведён к норме")
            if not run_in_hyperlink(run) and res.color(run, p):
                _set_black(run)
                self.done("Цвет текста изменён на чёрный")

    # ---- абзацы
    def fix_body_paragraph(self, p, role, prev_formula):
        n, res = self.n, self.res
        tol_i = n["first_line_indent_tolerance_cm"]
        caption = role.endswith("caption")
        text = p.text.strip()

        if not has_image(p):
            kind, v = res.line_spacing(p)
            size = self.chk._first_size(p)
            if caption:
                if kind == "ratio" and (v > n["caption_line_spacing_max"] + n["line_spacing_tolerance"]
                                        or v < 1.0 - n["line_spacing_tolerance"]):
                    p.paragraph_format.line_spacing = 1.0
                    self.done("Межстрочный интервал исправлен")
            else:
                bad = (abs(v - n["line_spacing"]) > n["line_spacing_tolerance"] if kind == "ratio"
                       else abs(v / (size * n["line_height_factor"]) - n["line_spacing"]) > 0.15)
                if bad:
                    p.paragraph_format.line_spacing = n["line_spacing"]
                    self.done("Межстрочный интервал исправлен")

        if role == "list":
            if BULLET_MARKERS.match(p.text):
                for t in _wts(p):
                    if (t.text or "").strip():
                        t.text = BULLET_MARKERS.sub("– ", (t.text or ""), count=1)
                        break
                self.done("Маркер списка заменён на дефис")
            return

        align = res.alignment(p)
        indent = res.first_line_indent_cm(p, self.chk._first_size(p))
        if role == "figure_caption":
            if align != "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.done("Подпись к рисунку выровнена по центру")
            if abs(indent) > tol_i:
                _set_first_line(p, 0)
                self.done("У подписи убран абзацный отступ")
            self.fix_caption_text(p, "Рисунок")
        elif role == "table_caption":
            if align not in ("left", "justify"):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self.done("Название таблицы выровнено по левому краю")
            if abs(indent) > tol_i:
                _set_first_line(p, 0)
                self.done("У названия таблицы убран абзацный отступ")
            self.fix_caption_text(p, "Таблица")
        else:
            continued = prev_formula and re.match(r"^\s*где\b", text, re.IGNORECASE)
            if align != "justify":
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                self.done("Выравнивание по ширине")
            if not continued and abs(indent - n["first_line_indent_cm"]) > tol_i:
                _set_first_line(p, n["first_line_indent_cm"])
                self.done("Абзацный отступ " + f"{n['first_line_indent_cm']:g}".replace(".", ",") + " см")

    def fix_caption_text(self, p, full):
        if _has_field(p):  # номер — поле SEQ: переписывать текст нельзя
            return
        new = _normalize_caption(p.text, full)
        if new is None or new == p.text.strip():
            return
        ts = _wts(p)
        if not ts:
            return
        ts[0].text = new
        ts[0].set(qn("xml:space"), "preserve")
        for t in ts[1:]:
            t.text = ""
        self.done("Подпись приведена к формату «" + full + " N – Название»")

    def fix_heading(self, p, role):
        res = self.res
        runs = self.chk._runs(p)
        ts = [t for t in _wts(p) if (t.text or "")]
        text = "".join((t.text or "") for t in ts)
        structural = role == "structural"

        if role == "heading" and runs and not all(res.bold(r, p) for r in runs):
            for r in runs:
                r.font.bold = True
            self.done("Заголовок раздела выделен жирным")
        if any(res.underline(r, p) for r in runs):
            for r in runs:
                r.font.underline = False
            self.done("Подчёркивание заголовка убрано")

        if structural:
            if any(c.islower() for c in text if c.isalpha()):
                for t in ts:
                    t.text = (t.text or "").upper()
                self.done("Заголовок структурного элемента набран прописными")
            if res.alignment(p) != "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_first_line(p, 0)
                self.done("Заголовок структурного элемента выровнен по центру")
        else:
            lead = len(text) - len(text.lstrip())
            m = CHAPTER_NUMBER_RE.match(text.lstrip())
            if m and m.group(2):
                self._delete_char(ts, lead + m.end(1))
                self.done("Убрана точка после номера раздела")
                ts = [t for t in _wts(p) if (t.text or "")]
                text = "".join((t.text or "") for t in ts)
            # первая буква после номера раздела: «2.1 структура» → «2.1 Структура»
            pos = len(HEADING_NUMBER_PREFIX_RE.match(text).group(0)) if HEADING_NUMBER_PREFIX_RE.match(text) else 0
            if text[pos:pos + 1].islower():
                for t in ts:
                    tt = t.text or ""
                    if pos < len(tt):
                        t.text = tt[:pos] + tt[pos].upper() + tt[pos + 1:]
                        break
                    pos -= len(tt)
                self.done("Заголовок начат с прописной буквы")

        if ends_with_sentence_dot(text):
            for t in reversed(ts):
                if (t.text or "").strip():
                    t.text = (t.text or "").rstrip()[:-1]
                    break
            self.done("Убрана точка в конце заголовка")

    @staticmethod
    def _delete_char(ts, index):
        pos = 0
        for t in ts:
            if pos + len((t.text or "")) > index:
                k = index - pos
                t.text = (t.text or "")[:k] + (t.text or "")[k + 1:]
                return
            pos += len((t.text or ""))

    def fix_toc(self, p):
        level = _toc_level(p)
        if level is None:
            return
        runs = [r for r in self.chk._runs(p) if not re.fullmatch(r"\d+", r.text.strip())]
        if not runs:
            return
        is_bold = all(self.res.bold(r, p) for r in runs)
        if level == 1 and not is_bold:
            for r in runs:
                r.font.bold = True
            self.done("Оглавление: заголовки разделов выделены жирным")
        elif level > 1 and is_bold:
            for r in runs:
                r.font.bold = False
            self.done("Оглавление: пункты сделаны не жирными")

    # ---- текст: пробелы
    def fix_text(self, p, role):
        changed = False
        for t in _wts(p):
            old = (t.text or "")
            new = MULTI_SPACE.sub(" ", old)
            if id(p._p) not in self._bib_ids:   # в записях списка « : » и « ; » — по ГОСТ Р 7.0.100
                new = SPACE_BEFORE_SIMPLE.sub(r"\1", new)
                new = SPACE_BEFORE_DOT.sub(r"\1", new)
                if role in ("body", "list") and not CODE_LIKE.search(old):
                    new = HYPHEN_AS_DASH.sub(" – ", new)
            if new != old:
                t.text = new
                changed = True
        ts = _wts(p)
        for a, b in zip(ts, ts[1:]):  # двойной пробел на стыке двух run'ов
            at, bt = a.text or "", b.text or ""
            if at.endswith(" ") and bt.startswith(" ") and bt.strip():
                b.text = bt.lstrip(" ")
                changed = True
        if changed:
            self.done("Исправлены лишние пробелы")
        if role != "heading" and role != "structural" and p.text[:1] in (" ", "\t"):
            for r in get_all_runs(p):
                for el in list(r._r):
                    if el.tag == qn("w:tab"):
                        r._r.remove(el)
                    elif el.tag == qn("w:t"):
                        if el.text and el.text.strip():
                            el.text = el.text.lstrip(" \t")
                            self.done("Убраны пробелы в начале абзаца")
                            return
                        el.text = ""
                    elif el.tag != qn("w:rPr"):
                        return

    # ---- список источников: только однозначные правки текста записи
    def fix_bibliography(self, items):
        """
        Word режет текст записи на много run'ов (проверка орфографии, гиперссылки), поэтому
        правки делаются по склеенному тексту абзаца и раскладываются обратно по run'ам
        (_replace_spans): форматирование и гиперссылки сохраняются.
        """
        _, entries = bib.find_section(self.chk, items)
        for p in entries:
            ts = _wts(p)
            full = "".join(t.text or "" for t in ts)
            if not full.strip():
                continue
            edits = []   # (начало, конец, новый текст) в координатах full

            m_num = bib.NUM_PREFIX.match(full)
            head_at = m_num.end() if m_num and not self.res.is_numbered(p) else 0
            if head_at and (m_num.group(1).startswith("[") or m_num.group(2) != "."):
                # «2)», «[2]», «2» → «2.» (ГОСТ 7.32-2017, п. 6.16)
                edits.append((m_num.start(), m_num.end(), f"{m_num.group(1).strip('[]')}. "))
            rest = full[head_at:]
            m = bib.HEAD_NO_COMMA.match(rest)
            if m:
                edits.append((head_at + m.start(), head_at + m.end(), f"{m.group(1)}, {bib._spaced(m.group(2))}"))
            else:
                m = re.match(r"(?:[А-ЯЁA-Z][а-яёa-z]+(?:-[А-ЯЁA-Z][а-яёa-z]+)?),\s+[А-ЯЁA-Z]\.(?=[А-ЯЁA-Z]\.)", rest)
                if m:
                    edits.append((head_at + m.end(), head_at + m.end(), " "))

            for mm in bib.OLD_ER.finditer(full):
                # «[Электронный ресурс]» в ГОСТ Р 7.0.100-2018 нет («Текст : электронный» можно не писать)
                start = mm.start()
                while start > 0 and full[start - 1] == " ":
                    start -= 1
                edits.append((start, mm.end(), ""))
            for mm in bib.HYPHEN_SEP.finditer(full):
                edits.append((mm.start(), mm.end(), " – "))
            for mm in bib.NO_SPACE_NUMBER.finditer(full):
                edits.append((mm.start(), mm.end(), "№ "))
            for mm in bib.CITY_ABBR.finditer(full):
                city = bib.CITY_FULL.get(mm.group(1).replace(" ", ""))
                if city:
                    edits.append((mm.start(), mm.end(), city + " :"))
            for mm in re.finditer(r"(?<=[–—]\s)([А-ЯЁ][а-яё-]+)(\s?):(\s?)(?=[«\"А-ЯЁA-Z])", full):
                if mm.group(2) != " " or mm.group(3) != " ":
                    edits.append((mm.start(), mm.end(), mm.group(1) + " : "))

            mu = bib.URL_LABEL.search(full)
            if mu:
                before = full[:mu.start()]
                stripped = before.rstrip()
                if stripped and not re.search(r"[–—]$", stripped):
                    sep = " – " if stripped.endswith(".") else ". – "
                    edits.append((len(stripped), mu.start(), sep))

            if _replace_spans(ts, edits):
                self.done("Исправлено оформление записей в списке источников")

    def move_misplaced_captions(self, items):
        """Название таблицы — над таблицей (п. 6.6.3), подпись рисунка — под рисунком (п. 6.5.7)."""
        moved = False
        for k, it in enumerate(items):
            if not it.get("misplaced"):
                continue
            step = -1 if it["role"] == "table_caption" else +1
            j = k + step
            while 0 <= j < len(items) and items[j]["kind"] == "p" and items[j]["role"] == "empty":
                j += step
            if not 0 <= j < len(items):
                continue
            target = items[j]["obj"]._element
            if step < 0:
                target.addprevious(it["obj"]._p)   # название — перед таблицей
                self.done("Название таблицы перенесено над таблицей")
            else:
                target.addnext(it["obj"]._p)       # подпись — после рисунка
                self.done("Подпись перенесена под рисунок")
            moved = True
        # подпись листинга под кодом → над кодом (если над кодом подписи нет)
        for block in self.chk.code_blocks(items):
            above = block[0]["i"] - 1
            while above >= 0 and items[above]["role"] == "empty":
                above -= 1
            below = block[-1]["i"] + 1
            while below < len(items) and items[below]["role"] == "empty":
                below += 1
            if below < len(items) and items[below]["role"] == "listing_caption" \
                    and not (above >= 0 and items[above]["role"] == "listing_caption"):
                block[0]["obj"]._p.addprevious(items[below]["obj"]._p)
                self.done("Подпись листинга перенесена над кодом")
                moved = True
        return moved

    def remove_extra_empty(self, items):
        prev_empty = False
        for it in items:
            if it["kind"] != "p" or it["role"] != "empty":
                prev_empty = False
                continue
            p = it["obj"]
            blocked = has_page_break(p) or p._p.xpath(".//w:bookmarkStart | .//w:fldChar | .//w:fldSimple")
            if blocked:
                prev_empty = False
                continue
            if prev_empty:
                p._p.getparent().remove(p._p)
                self.done("Удалены лишние пустые абзацы")
            prev_empty = True

    # ---- таблицы
    def fix_table(self, table):
        n, res = self.n, self.res
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                for p in cell.paragraphs:
                    for run in self.chk._runs(p):
                        fname = res.font_name(run, p)
                        if fname and fname != n["font_name"] and fname.lower() not in ALT_FONTS:
                            _set_font_name(run, n["font_name"])
                            self.done("Шрифт заменён на " + n["font_name"])


def fix_docx(src_path, dst_path, profile, before=None):
    """
    Исправляет src_path → dst_path по профилю норм. Возвращает FixResult
    (что исправлено + замечания до и после).

    before — уже готовый отчёт проверки исходного файла (например, из кеша
    после /api/check): тогда повторная проверка исходника не нужна.
    """
    if before is None:
        before = Report(profile_key=profile["key"], profile_title=profile["title"])
        DocxChecker(Document(src_path), before, profile).run()

    doc = Document(src_path)
    fx = _Fixer(doc, profile)
    fx.run()
    # отметки «здесь была граница страницы» устарели — Word пересоздаст их сам
    for el in list(doc.element.body.iter(qn("w:lastRenderedPageBreak"))):
        el.getparent().remove(el)
    doc.save(dst_path)

    after = Report(profile_key=profile["key"], profile_title=profile["title"])
    DocxChecker(Document(dst_path), after, profile).run()
    return FixResult(applied=dict(fx.applied), before=before, after=after)
