# -*- coding: utf-8 -*-
"""Тесты автоисправления: исправляем нарушения, не ломаем корректное, не меняем текст."""

import hashlib
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, Cm, RGBColor

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from checker import run_normocontrol  # noqa: E402
from checker.fixers import fix_docx  # noqa: E402
from checker.norms import get_profile  # noqa: E402
import make_docs as md  # noqa: E402


def fix(doc, tmp_path, profile="vkr"):
    src, dst = tmp_path / "src.docx", tmp_path / "fixed.docx"
    doc.save(src)
    return src, dst, fix_docx(str(src), str(dst), get_profile(profile))


def words(path):
    return re.sub(r"\s+", " ", " ".join(p.text for p in Document(str(path)).paragraphs)).strip()


def broken_document():
    doc = md.good_document()
    md.set_style_font(doc.styles["Normal"], "Arial", 12)
    pf = doc.styles["Normal"].paragraph_format
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Cm(0.5)
    doc.sections[0].left_margin = Cm(2.0)
    doc.sections[0].right_margin = Cm(3.0)
    doc.sections[0].footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    p.add_run("Красный текст с  двойным пробелом и лишним пробелом , в предложении.").font.color.rgb = RGBColor(255, 0, 0)
    doc.add_paragraph("• маркированный пункт")
    h = doc.add_heading("2. Практическая часть.", level=1)
    doc.styles["Heading 1"].font.bold = False
    doc.styles["Heading 1"].font.color.rgb = RGBColor(0x2F, 0x54, 0x96)
    s = md.structural(doc, "Приложение")
    s.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return doc


def test_good_document_untouched(tmp_path):
    _, dst, r = fix(md.good_document(), tmp_path)
    assert r.applied == {}
    assert r.after.is_clean()


def test_broken_document_gets_clean(tmp_path):
    _, dst, r = fix(broken_document(), tmp_path)
    assert r.before.issues and r.total_applied > 0
    rep = run_normocontrol(str(dst), "vkr")
    assert rep.is_clean(), "\n".join(f"[{i.category}] {i.location}: {i.message}" for i in rep.issues)


def test_fix_is_idempotent(tmp_path):
    _, dst, r = fix(broken_document(), tmp_path)
    r2 = fix_docx(str(dst), str(tmp_path / "fixed2.docx"), get_profile("vkr"))
    assert r2.applied == {}


def test_source_file_not_modified(tmp_path):
    doc = broken_document()
    src = tmp_path / "src.docx"
    doc.save(src)
    before = hashlib.sha256(src.read_bytes()).hexdigest()
    fix_docx(str(src), str(tmp_path / "out.docx"), get_profile("vkr"))
    assert hashlib.sha256(src.read_bytes()).hexdigest() == before


def test_text_preserved_except_known_edits(tmp_path):
    src, dst, _ = fix(broken_document(), tmp_path)
    norm = lambda s: re.sub(r"[\s.,•–]+", "", s).lower()
    assert norm(words(src)) == norm(words(dst))


def test_heading_with_own_page_break_gets_no_extra_break(tmp_path):
    # разрыв страницы внутри самого заголовка: pageBreakBefore дал бы пустую страницу
    doc = md.good_document()
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    h = doc.add_heading("", level=1)
    h.add_run().add_break(WD_BREAK.PAGE)
    h.add_run("3 Новый раздел")
    _, dst, r = fix(doc, tmp_path)
    # единственное исправление — у «1 Теоретические основы» (стиль без разрыва); у «3 Новый раздел» разрыв уже есть
    assert r.applied.get("Заголовок начат с новой страницы") == 1


def test_missing_page_break_is_added(tmp_path):
    doc = md.good_document()
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    doc.add_heading("3 Раздел без разрыва", level=1)
    _, dst, r = fix(doc, tmp_path)
    assert r.applied.get("Заголовок начат с новой страницы") == 2  # «1 Теоретические основы» и «3 Раздел»
    assert r.after.is_clean()


def test_abbreviation_dot_kept(tmp_path):
    doc = md.good_document()
    img = doc.add_paragraph()
    img.add_run().add_picture(md._png(), width=Cm(3))
    cap = doc.add_paragraph("Рисунок 2 – Затраты, тыс. руб.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing = 1.0
    _, dst, r = fix(doc, tmp_path)
    assert any(p.text == "Рисунок 2 – Затраты, тыс. руб." for p in Document(str(dst)).paragraphs)
    assert r.after.is_clean()


def test_captions_normalized(tmp_path):
    doc = md.good_document()
    img = doc.add_paragraph()
    img.add_run().add_picture(md._png(), width=Cm(3))
    cap = doc.add_paragraph("Рис. 2. схема системы.")
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _, dst, r = fix(doc, tmp_path)
    assert any(p.text == "Рисунок 2 – Схема системы" for p in Document(str(dst)).paragraphs)


def test_page_numbers_added(tmp_path):
    doc = md.new_document(page_numbers=False)
    md.structural(doc, "ВВЕДЕНИЕ", page_break=False)
    md.body(doc)
    _, dst, r = fix(doc, tmp_path)
    assert "Нумерация страниц" not in {i.category for i in r.after.issues}


def test_field_caption_not_rewritten(tmp_path):
    # у подписи с полем SEQ переписывать текст нельзя — номер превратится в мёртвый текст
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = md.good_document()
    img = doc.add_paragraph()
    img.add_run().add_picture(md._png(), width=Cm(3))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.add_run("Рисунок ")
    r = cap.add_run()
    for kind, txt in (("begin", None), ("instr", " SEQ Рисунок "), ("end", None)):
        el = OxmlElement("w:instrText" if kind == "instr" else "w:fldChar")
        if kind == "instr":
            el.text = txt
        else:
            el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
    cap.add_run("2. подпись")
    _, dst, _ = fix(doc, tmp_path)
    assert "SEQ" in Document(str(dst)).element.xml
