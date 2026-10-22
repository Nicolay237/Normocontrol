# -*- coding: utf-8 -*-
"""Листинги программного кода: распознавание, проверки по правилам методичек, автоисправление."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import make_docs as md
from checker import run_normocontrol
from checker.fixers import fix_docx
from checker.norms import PROFILES

CODE = [
    "vector<int> Dijkstra(int start)",
    "{",
    "    vector<int> dist(V, INF);",
    "    dist[start] = 0;",
    "",
    "    return dist;",
    "}",
]


def code_lines(doc, lines=CODE, font="Courier New", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0.0):
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.first_line_indent = Cm(indent)
        p.paragraph_format.line_spacing = 1.0
        if line:
            r = p.add_run(line)
            r.font.name, r.font.size = font, Pt(size)


def caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0


def work(**code_kw):
    doc = md.new_document()
    md.structural(doc, "ВВЕДЕНИЕ", page_break=False)
    md.body(doc)
    md.body(doc, "Алгоритм Дейкстры реализован так (листинг 1) [1].")
    caption(doc, code_kw.pop("caption", "Листинг 1 – Алгоритм Дейкстры"))
    code_lines(doc, **code_kw)
    md.body(doc)
    md.structural(doc, "ЗАКЛЮЧЕНИЕ")
    md.body(doc)
    md.structural(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    md.body(doc, "1. Иванов, И. И. Алгоритмы на графах / И. И. Иванов. – Москва : Наука, 2020. – 240 с.")
    return doc


def check(doc, tmp_path, profile="vkr"):
    path = tmp_path / "w.docx"
    doc.save(path)
    return run_normocontrol(str(path), profile)


def test_proper_listing_has_no_issues(tmp_path):
    # Courier 12 пт, одинарный, по левому краю, двойные пробелы и пустая строка внутри, arr[...] в коде
    assert check(work(), tmp_path).issues == []


def test_listing_formatting_errors(tmp_path):
    msgs = [i.message for i in check(work(align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25, size=16,
                                          caption="Введение в алгоритм."), tmp_path).issues
            if i.category == "Листинги"]
    joined = "\n".join(msgs)
    assert "по левому краю" in joined and "абзацного отступа" in joined and "кегль листинга 16" in joined
    assert "нет подписи" in joined


def test_code_in_times_is_detected(tmp_path):
    report = check(work(font="Times New Roman", size=14), tmp_path)
    assert any("моноширинный шрифт" in i.message for i in report.issues)
    # строки кода не проверяются как основной текст (красная строка, выравнивание по ширине, пробелы)
    assert not [i for i in report.issues if i.category != "Листинги"
                and ("Dijkstra" in i.location or "dist" in i.location)]
    assert not any("нет подписи" in i.message for i in report.issues)   # подпись над кодом найдена


def test_figure_caption_under_code_is_accepted(tmp_path):
    doc = md.new_document()
    md.structural(doc, "ВВЕДЕНИЕ", page_break=False)
    md.body(doc)
    code_lines(doc)
    p = doc.add_paragraph("Рисунок 1 – Алгоритм Дейкстры")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    assert not [i for i in check(doc, tmp_path).issues if i.category in ("Листинги", "Нумерация")]


def test_long_listing_belongs_in_appendix(tmp_path):
    lines = [f"x{i} = f({i});" for i in range(65)]
    assert any("в приложение" in i.message for i in check(work(lines=lines), tmp_path).issues)


def test_bibliography_with_urls_is_not_code(tmp_path):
    doc = md.good_document()
    for i in range(2, 6):
        md.body(doc, f"{i}. Face API // Microsoft. – URL: https://learn.microsoft.com/a{i} (дата обращения: 01.03.2025).")
    report = check(doc, tmp_path)
    assert not [i for i in report.issues if i.category == "Листинги"]


def test_autofix_keeps_code_and_aligns_left(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    work(align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25).save(src)
    result = fix_docx(str(src), str(dst), PROFILES["vkr"])
    assert not [i for i in result.after.issues if i.category == "Листинги"]
    code = [p for p in Document(dst).paragraphs if "dist" in p.text or "Dijkstra" in p.text]
    for p in code:
        assert p.alignment == WD_ALIGN_PARAGRAPH.LEFT
        assert p.runs[0].font.name == "Courier New" and p.runs[0].font.size == Pt(12)   # шрифт не тронут
        assert p.paragraph_format.line_spacing == 1.0                                    # интервал не тронут
    # пустая строка внутри кода не удалена как «лишняя»
    texts = [p.text for p in Document(dst).paragraphs]
    i = texts.index("    dist[start] = 0;")
    assert texts[i + 1] == "" and texts[i + 2] == "    return dist;"
