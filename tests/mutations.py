# -*- coding: utf-8 -*-
"""
Каталог ошибок оформления для проверки сервиса «от противного».

Каждая ошибка (мутация) портит одно место эталонной курсовой (scripts/make_sample_coursework.py),
и проверка обязана её поймать: в отчёте должно появиться замечание, содержащее `expect`.
`fixable` — должно ли автоисправление убрать это замечание.

Тот же каталог собирает «работу со всеми ошибками сразу» — scripts/make_broken_coursework.py.
"""

import copy
from dataclasses import dataclass, field
from typing import Callable

from docx.enum.section import WD_ORIENT  # noqa: F401  (на будущее: альбомные листы)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


@dataclass
class Mutation:
    key: str
    title: str                  # что испорчено — для отчёта
    gost: str                   # какое требование нарушено
    apply: Callable
    expect: str                 # фрагмент ожидаемого замечания
    fixable: bool = False       # должно ли автоисправление его убрать
    profile: str = "vkr"
    standalone: bool = False    # ломает слишком много — в «работу со всеми ошибками» не берём
    tags: list = field(default_factory=list)


# ── поиск мест в эталоне ──────────────────────────────────────────────────

def para(doc, start):
    for p in doc.paragraphs:
        if p.text.lstrip().startswith(start):
            return p
    raise LookupError(start)


def heading(doc, start):
    """Заголовок раздела (а не одноимённая строка содержания)."""
    for p in doc.paragraphs:
        if p.text.startswith(start) and p.style.name.startswith("Heading"):
            return p
    raise LookupError(start)


def body_para(doc):
    return para(doc, "Поиск кратчайшего пути")


def body_para2(doc):
    return para(doc, "Актуальность работы связана")


def runs(p, **kw):
    for r in p.runs:
        for k, v in kw.items():
            if k == "font":
                r.font.name = v
                rpr = r._r.get_or_add_rPr()
                rf = rpr.find(qn("w:rFonts"))
                if rf is None:
                    rf = rpr.makeelement(qn("w:rFonts"), {})
                    rpr.append(rf)
                for a in ("w:ascii", "w:hAnsi", "w:cs"):
                    rf.set(qn(a), v)
            elif k == "size":
                r.font.size = Pt(v)
            elif k == "color":
                r.font.color.rgb = RGBColor.from_string(v)
            elif k == "bold":
                r.bold = v
            elif k == "underline":
                r.underline = v


def set_text(p, text):
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = text


def replace_text(p, old, new):
    full = p.text
    assert old in full, (old, full)
    set_text(p, full.replace(old, new, 1))


def source(doc, n):
    return para(doc, f"{n}. ")


def table_of(doc):
    return doc.tables[0]


# ── мутации ───────────────────────────────────────────────────────────────

def m_margin(side, cm):
    def f(doc):
        setattr(doc.sections[0], f"{side}_margin", Cm(cm))
    return f


def m_letter(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)


def m_no_page_numbers(doc):
    s = doc.sections[0]
    for p in s.footer.paragraphs:
        p._p.getparent().remove(p._p)
    s.footer.add_paragraph()


def m_page_number_right(doc):
    for p in doc.sections[0].footer.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def m_page_number_top(doc):
    s = doc.sections[0]
    footer_p = s.footer.paragraphs[0]
    header_p = s.header.paragraphs[0]
    for r in list(footer_p.runs):
        header_p._p.append(r._r)
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def m_heading_dot(doc):
    p = heading(doc, "1.1 Задача поиска")
    replace_text(p, "кратчайшего пути", "кратчайшего пути.")


def m_heading_number_dot(doc):
    replace_text(heading(doc, "1.2 Обзор алгоритмов"), "1.2 ", "1.2. ")


def m_heading_not_bold(doc):
    runs(heading(doc, "2.1 Структура данных"), bold=False)


def m_heading_underline(doc):
    runs(heading(doc, "2.1 Структура данных"), underline=True)


def m_heading_lowercase(doc):
    replace_text(heading(doc, "2.1 Структура данных"), "Структура", "структура")


def m_heading_size(doc):
    runs(heading(doc, "2.2 Реализация алгоритма"), size=18)


def m_heading_color(doc):
    runs(heading(doc, "2.2 Реализация алгоритма"), color="1F4E79")


def m_heading_no_page_break(doc):
    p = heading(doc, "3 Тестирование")
    p.paragraph_format.page_break_before = False


def m_structural_lowercase(doc):
    replace_text(para(doc, "ЗАКЛЮЧЕНИЕ"), "ЗАКЛЮЧЕНИЕ", "Заключение")


def m_structural_left(doc):
    p = para(doc, "ЗАКЛЮЧЕНИЕ")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def m_structural_dot(doc):
    replace_text(para(doc, "ЗАКЛЮЧЕНИЕ"), "ЗАКЛЮЧЕНИЕ", "ЗАКЛЮЧЕНИЕ.")


def m_no_conclusion(doc):
    p = para(doc, "ЗАКЛЮЧЕНИЕ")
    p._p.getparent().remove(p._p)


def m_no_toc(doc):
    for p in list(doc.paragraphs):
        if p.text == "СОДЕРЖАНИЕ" or p.style.name.startswith("toc"):
            p._p.getparent().remove(p._p)


def m_toc_sub_bold(doc):
    for p in doc.paragraphs:
        if p.style.name == "toc 2":
            p.runs[0].bold = True


def m_toc_main_not_bold(doc):
    for p in doc.paragraphs:
        if p.style.name == "toc 1":
            p.runs[0].bold = False


def m_font_arial(doc):
    runs(body_para(doc), font="Arial")


def m_font_size_12(doc):
    runs(body_para(doc), size=12)


def m_font_size_16(doc):
    runs(body_para(doc), size=16)


def m_font_color(doc):
    runs(body_para(doc), color="C00000")


def m_spacing_single(doc):
    body_para(doc).paragraph_format.line_spacing = 1.0


def m_spacing_double(doc):
    body_para(doc).paragraph_format.line_spacing = 2.0


def m_indent_zero(doc):
    body_para(doc).paragraph_format.first_line_indent = Cm(0)


def m_indent_large(doc):
    body_para(doc).paragraph_format.first_line_indent = Cm(2)


def m_indent_spaces(doc):
    p = body_para2(doc)
    p.paragraph_format.first_line_indent = Cm(0)
    p.runs[0].text = "     " + p.runs[0].text


def m_align_left(doc):
    body_para(doc).alignment = WD_ALIGN_PARAGRAPH.LEFT


def m_align_center(doc):
    body_para(doc).alignment = WD_ALIGN_PARAGRAPH.CENTER


def m_double_space(doc):
    replace_text(body_para(doc), "одна из базовых", "одна  из базовых")


def m_space_before_comma(doc):
    replace_text(body_para(doc), "сервисы, логистические", "сервисы , логистические")


def m_empty_lines(doc):
    p = body_para(doc)
    for _ in range(3):
        p._p.addnext(copy.deepcopy(p._p))
    for el in list(p._p.itersiblings())[:3]:
        for r in el.findall(qn("w:r")):
            el.remove(r)


def m_bullet_marker(doc):
    p = para(doc, "- изучить постановку")
    replace_text(p, "- изучить", "• изучить")


def m_fig_short_word(doc):
    replace_text(para(doc, "Рисунок 1 – "), "Рисунок 1", "Рис. 1")


def m_fig_dot(doc):
    replace_text(para(doc, "Рисунок 1 – "), "дорожной сети", "дорожной сети.")


def m_fig_left(doc):
    para(doc, "Рисунок 1 – ").alignment = WD_ALIGN_PARAGRAPH.LEFT


def m_fig_no_dash(doc):
    replace_text(para(doc, "Рисунок 1 – "), "Рисунок 1 – ", "Рисунок 1. ")


def m_fig_lowercase(doc):
    replace_text(para(doc, "Рисунок 1 – "), "Взвешенный", "взвешенный")


def m_fig_numbering(doc):
    replace_text(para(doc, "Рисунок 1 – "), "Рисунок 1", "Рисунок 2")


def m_fig_removed(doc):
    p = para(doc, "Рисунок 1 – ")
    img = p._p.getprevious()
    img.getparent().remove(img)


def m_fig_caption_above(doc):
    p = para(doc, "Рисунок 1 – ")
    img = p._p.getprevious()
    img.addprevious(p._p)


def m_tab_center(doc):
    para(doc, "Таблица 1 – ").alignment = WD_ALIGN_PARAGRAPH.CENTER


def m_tab_indent(doc):
    para(doc, "Таблица 1 – ").paragraph_format.first_line_indent = Cm(1.25)


def m_tab_below(doc):
    p = para(doc, "Таблица 1 – ")
    table_of(doc)._tbl.addnext(p._p)


def m_tab_removed(doc):
    t = table_of(doc)._tbl
    t.getparent().remove(t)


def m_tab_dot(doc):
    replace_text(para(doc, "Таблица 1 – "), "кратчайших путей", "кратчайших путей.")


def m_tab_cell_font(doc):
    for cell in table_of(doc).rows[1].cells:
        runs(cell.paragraphs[0], font="Arial")


def m_code_times(doc):
    for p in doc.paragraphs:
        if any(r.font.name == "Courier New" for r in p.runs):
            for r in p.runs:
                r.font.name = "Times New Roman"


def m_code_justify(doc):
    start = para(doc, "Листинг 1 – ")
    el = start._p.getnext()
    from docx.text.paragraph import Paragraph
    for _ in range(25):
        Paragraph(el, start._parent).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        el = el.getnext()


def m_listing_no_caption(doc):
    p = para(doc, "Листинг 1 – ")
    p._p.getparent().remove(p._p)


def m_bib_no_comma(doc):
    replace_text(source(doc, 2), "Седжвик, Р. Алгоритмы", "Седжвик Р. Алгоритмы")


def m_bib_city_abbr(doc):
    replace_text(source(doc, 2), "Москва : Вильямс", "М. : Вильямс")


def m_bib_electronic(doc):
    replace_text(source(doc, 6), "Документация по стандартной библиотеке C++.",
                 "Документация по стандартной библиотеке C++ [Электронный ресурс].")


def m_bib_no_url_label(doc):
    replace_text(source(doc, 6), "URL: ", "")


def m_bib_no_access_date(doc):
    replace_text(source(doc, 6), " (дата обращения: 10.03.2026)", "")


def m_bib_hyphen(doc):
    replace_text(source(doc, 3), "2004. – 368 с.", "2004. - 368 с.")


def m_bib_no_volume(doc):
    replace_text(source(doc, 3), ". – 368 с.", "")


def m_bib_numbering(doc):
    replace_text(source(doc, 2), "2. ", "2) ")


def m_bib_duplicate(doc):
    replace_text(source(doc, 6), "6. Документация", "7. Документация")
    src = source(doc, 5)
    dup = copy.deepcopy(src._p)
    src._p.addnext(dup)
    from docx.text.paragraph import Paragraph
    replace_text(Paragraph(dup, src._parent), "5. ", "6. ")


def m_cite_missing_source(doc):
    replace_text(para(doc, "По результатам сравнения"), "[4]", "[9]")


def m_cite_order(doc):
    replace_text(para(doc, "Дорожную сеть удобно"), "[1]", "[3]")


def m_uncited_source(doc):
    replace_text(para(doc, "Для хранения списков используется"), " [6]", "")


def m_appendix_title_dot(doc):
    replace_text(para(doc, "Полный текст программы"), "программы", "программы.")


def m_font_11(doc):
    runs(body_para(doc), size=11)


def m_caption_small(doc):
    runs(para(doc, "Рисунок 1 – "), size=10)


def m_listing_caption_below(doc):
    cap = para(doc, "Листинг 1 – ")
    el = cap._p
    last = el
    for _ in range(len([1 for _ in range(25)])):
        last = last.getnext()
    last.addnext(el)


def m_bib_tight_initials(doc):
    replace_text(source(doc, 2), "Седжвик, Р. Алгоритмы", "Седжвик, Р.А. Алгоритмы")


def m_bib_no_space_number(doc):
    replace_text(source(doc, 4), "№ 4", "№4")


def m_struct_no_break(doc):
    para(doc, "ЗАКЛЮЧЕНИЕ").paragraph_format.page_break_before = False


def m_tab_short_word(doc):
    replace_text(para(doc, "Таблица 1 – "), "Таблица 1", "Табл. 1")


def m_hyphen_dash(doc):
    replace_text(body_para2(doc), "связана с тем, что", "связана с тем - что")


def m_list_marker_star(doc):
    replace_text(para(doc, "- выбрать структуру"), "- выбрать", "* выбрать")


def m_normal_calibri(doc):
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    rf = st.element.get_or_add_rPr().find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Calibri")
    st.paragraph_format.line_spacing = 1.15


# Ошибки для «работы со всеми ошибками сразу»: в одном документе, не мешая друг другу
# (одно свойство одного места портится один раз; порядок важен — сначала то, что ищет текст)
COMBINED = [
    "margin_left", "margin_top", "page_number_right",
    "font_arial", "font_16", "font_color", "spacing_single", "indent_zero", "align_left",
    "double_space", "space_before_comma", "indent_spaces", "hyphen_dash", "empty_lines", "bullet", "list_star",
    "heading_dot", "heading_number_dot", "heading_not_bold", "heading_underline", "heading_lowercase",
    "heading_size", "heading_color", "heading_no_break",
    "struct_left", "struct_no_break", "struct_lowercase", "toc_sub_bold",
    "fig_dot", "fig_left", "fig_lowercase", "caption_small",
    "tab_indent", "tab_dot", "tab_cell_font", "tab_below",
    "code_justify", "listing_caption_below",
    "bib_city", "bib_no_comma", "bib_er", "bib_access_date", "bib_hyphen", "bib_number_sign",
    "cite_missing", "cite_order", "uncited", "bib_numbering",
    "appendix_dot",
]

MUTATIONS = [
    # страница (ГОСТ 7.32-2017, 6.1.1; 6.3.1)
    Mutation("margin_left", "Левое поле 2,5 см", "6.1.1: левое 30 мм", m_margin("left", 2.5), "левое поле", True),
    Mutation("margin_right", "Правое поле 2,5 см", "6.1.1: правое 15 мм", m_margin("right", 2.5), "правое поле", True),
    Mutation("margin_top", "Верхнее поле 1 см", "6.1.1: верхнее 20 мм", m_margin("top", 1), "верхнее поле", True),
    Mutation("margin_bottom", "Нижнее поле 3 см", "6.1.1: нижнее 20 мм", m_margin("bottom", 3), "нижнее поле", True),
    Mutation("page_letter", "Формат Letter вместо A4", "6.1.1: формат A4", m_letter, "ожидается A4", True),
    Mutation("no_page_numbers", "Нет нумерации страниц", "6.3.1", m_no_page_numbers, "номера страниц не найдены", True),
    Mutation("page_number_right", "Номер страницы справа", "6.3.1: по центру", m_page_number_right, "выровнен справа", True),
    Mutation("page_number_top", "Номер страницы вверху", "6.3.1: внизу", m_page_number_top, "верхнем колонтитуле"),
    # основной текст (6.1.1)
    Mutation("font_arial", "Абзац набран Arial", "6.1.1: Times New Roman", m_font_arial, "шрифт «Arial»", True),
    Mutation("font_12_vkr", "Кегль 12 в ВКР", "методички ВКР: 14 пт", m_font_size_12, "размер шрифта 12", True),
    Mutation("font_16", "Кегль 16", "6.1.1", m_font_size_16, "размер шрифта 16", True),
    Mutation("font_color", "Красный текст", "6.1.1: чёрный цвет", m_font_color, "цвет текста", True),
    Mutation("spacing_single", "Одинарный интервал", "6.1.1: полуторный", m_spacing_single, "интервал 1", True),
    Mutation("spacing_double", "Двойной интервал", "6.1.1: полуторный", m_spacing_double, "интервал 2", True),
    Mutation("indent_zero", "Нет красной строки", "6.1.1: отступ 1,25 см", m_indent_zero, "отступ первой строки 0", True),
    Mutation("indent_2cm", "Красная строка 2 см", "6.1.1: 1,25 см", m_indent_large, "отступ первой строки 2", True),
    Mutation("indent_spaces", "Красная строка пробелами", "6.1.1", m_indent_spaces, "начинается с пробела", True),
    Mutation("align_left", "Текст по левому краю", "выравнивание по ширине", m_align_left, "по левому краю", True),
    Mutation("align_center", "Текст по центру", "выравнивание по ширине", m_align_center, "по центру, ожидается", True),
    Mutation("double_space", "Двойной пробел", "типографика", m_double_space, "двойной", True),
    Mutation("space_before_comma", "Пробел перед запятой", "типографика", m_space_before_comma, "пробел перед знаком", True),
    Mutation("empty_lines", "Три пустые строки подряд", "типографика", m_empty_lines, "пустых абзаца", True),
    Mutation("bullet", "Перечисление с «•»", "6.4.5: дефис", m_bullet_marker, "маркер перечисления", True),
    # заголовки (6.2, 6.4.1)
    Mutation("heading_dot", "Точка в конце заголовка", "6.2.3", m_heading_dot, "точки"),
    Mutation("heading_number_dot", "Точка после номера «1.2.»", "6.4.1", m_heading_number_dot, "после номера раздела", True),
    Mutation("heading_not_bold", "Заголовок не жирный", "6.2.3: полужирный", m_heading_not_bold, "жирным", True),
    Mutation("heading_underline", "Заголовок подчёркнут", "6.2.3", m_heading_underline, "подчёркнут", True),
    Mutation("heading_lowercase", "Заголовок со строчной", "6.2.3", m_heading_lowercase, "с прописной", True),
    Mutation("heading_size", "Заголовок 18 пт", "6.1.1", m_heading_size, "размер шрифта заголовка", True),
    Mutation("heading_color", "Синий заголовок", "6.1.1: чёрный", m_heading_color, "цвет заголовка", True),
    Mutation("heading_no_break", "Раздел не с новой страницы", "6.2.1", m_heading_no_page_break, "с новой страницы", True),
    Mutation("struct_lowercase", "«Заключение» строчными", "6.2.1: прописными", m_structural_lowercase, "ПРОПИСНЫМИ", True),
    Mutation("struct_left", "«ЗАКЛЮЧЕНИЕ» слева", "6.2.1: по центру", m_structural_left, "по центру", True),
    Mutation("struct_dot", "Точка после «ЗАКЛЮЧЕНИЕ.»", "6.2.1: без точки", m_structural_dot, "точки"),
    Mutation("no_conclusion", "Нет заключения", "4.1: обязательный элемент", m_no_conclusion, "«ЗАКЛЮЧЕНИЕ»", standalone=True),
    Mutation("no_toc", "Нет содержания", "4.1: обязательный элемент", m_no_toc, "«СОДЕРЖАНИЕ»", standalone=True),
    Mutation("toc_sub_bold", "Подразделы в содержании жирные", "методички", m_toc_sub_bold, "не должен быть жирным", True),
    Mutation("toc_main_not_bold", "Разделы в содержании не жирные", "методички", m_toc_main_not_bold, "выделен жирным", True),
    # рисунки (6.5)
    Mutation("fig_short", "«Рис. 1»", "6.5.7: слово полностью", m_fig_short_word, "полностью", True),
    Mutation("fig_dot", "Точка в конце подписи", "6.5.7", m_fig_dot, "точка не ставится", True),
    Mutation("fig_left", "Подпись рисунка слева", "6.5.7: по центру", m_fig_left, "по центру", True),
    Mutation("fig_no_dash", "«Рисунок 1. Название»", "6.5.7: через тире", m_fig_no_dash, "пробел, тире, пробел", True),
    Mutation("fig_lowercase", "Подпись со строчной", "6.5.7", m_fig_lowercase, "с прописной", True),
    Mutation("fig_numbering", "Первый рисунок — «Рисунок 2»", "6.5.5: сквозная нумерация", m_fig_numbering, "сквозная нумерация"),
    Mutation("fig_removed", "Подпись без рисунка", "6.5", m_fig_removed, "нет рисунка", standalone=True),
    Mutation("fig_above", "Подпись над рисунком", "6.5.7: под рисунком", m_fig_caption_above, "над рисунком", True),
    # таблицы (6.6)
    Mutation("tab_center", "Название таблицы по центру", "6.6.3: слева", m_tab_center, "слева над таблицей", True),
    Mutation("tab_indent", "Название таблицы с отступом", "6.6.3: без отступа", m_tab_indent, "без отступа", True),
    Mutation("tab_below", "Название под таблицей", "6.6.3: над таблицей", m_tab_below, "под таблицей", True),
    Mutation("tab_removed", "Название без таблицы", "6.6", m_tab_removed, "нет самой таблицы", standalone=True),
    Mutation("tab_dot", "Точка в конце названия таблицы", "6.6.3", m_tab_dot, "точка не ставится", True),
    Mutation("tab_cell_font", "Arial в таблице", "6.1.1", m_tab_cell_font, "шрифт «Arial»", True),
    # листинги (методички вузов)
    Mutation("code_times", "Код набран Times", "моноширинный шрифт", m_code_times, "моноширинный"),
    Mutation("code_justify", "Код по ширине", "по левому краю", m_code_justify, "по левому краю, а не", True),
    Mutation("listing_no_caption", "Листинг без подписи", "подпись над кодом", m_listing_no_caption, "нет подписи"),
    # список источников (ГОСТ Р 7.0.100-2018, ГОСТ 7.32-2017 6.16, ГОСТ Р 7.0.5-2008)
    Mutation("bib_no_comma", "«Седжвик Р.» без запятой", "7.0.100", m_bib_no_comma, "запятая", True),
    Mutation("bib_city", "«М. :» вместо «Москва :»", "7.0.100: город полностью", m_bib_city_abbr, "место издания", True),
    Mutation("bib_er", "«[Электронный ресурс]»", "7.0.100", m_bib_electronic, "Электронный ресурс", True),
    Mutation("bib_url_label", "Ссылка без «URL:»", "7.0.100", m_bib_no_url_label, "URL:"),
    Mutation("bib_access_date", "Нет даты обращения", "7.0.100", m_bib_no_access_date, "дата обращения"),
    Mutation("bib_hyphen", "Дефис вместо тире", "7.0.100", m_bib_hyphen, "тире", True),
    Mutation("bib_volume", "Нет объёма книги", "7.0.100", m_bib_no_volume, "объём"),
    Mutation("bib_numbering", "Номер «2)»", "6.16: цифра с точкой", m_bib_numbering, "цифра с точкой", True),
    Mutation("bib_duplicate", "Источник повторяется", "6.16", m_bib_duplicate, "повторяется"),
    Mutation("cite_missing", "Ссылка [9] при 6 источниках", "7.0.5", m_cite_missing_source, "несуществующий"),
    Mutation("cite_order", "Первая ссылка — [3]", "6.16: по порядку упоминания", m_cite_order, "порядке упоминания"),
    Mutation("uncited", "Источник без ссылки (ВКР)", "6.16", m_uncited_source, "нет ссылки"),
    # приложения (6.17)
    Mutation("appendix_dot", "Точка в заголовке приложения", "6.17", m_appendix_title_dot, "точки", True),
    # ещё
    Mutation("font_11_course", "Кегль 11 в курсовой", "6.1.1: не менее 12 пт", m_font_11, "размер шрифта 11", True,
             profile="coursework"),
    Mutation("caption_small", "Подпись рисунка 10 пт", "6.1.1: не менее 12 пт", m_caption_small, "размер шрифта 10", True),
    Mutation("listing_caption_below", "Подпись листинга под кодом", "подпись над кодом", m_listing_caption_below,
             "ставится над кодом", True),
    Mutation("bib_tight", "Инициалы «Р.А.» без пробела", "7.0.100", m_bib_tight_initials, "через пробел", True),
    Mutation("bib_number_sign", "«№4» без пробела", "7.0.100", m_bib_no_space_number, "№", True),
    Mutation("struct_no_break", "Заключение не с новой страницы", "6.2.1", m_struct_no_break, "с новой страницы", True),
    Mutation("tab_short", "«Табл. 1»", "6.6.3: слово полностью", m_tab_short_word, "полностью", True),
    Mutation("hyphen_dash", "Дефис вместо тире в тексте", "типографика", m_hyphen_dash, "ставится тире", True),
    Mutation("list_star", "Перечисление со «*»", "6.4.5: дефис", m_list_marker_star, "маркер перечисления", True),
    Mutation("normal_calibri", "Весь текст Calibri 11, интервал 1,15", "6.1.1", m_normal_calibri, "шрифт «Calibri»", True,
             standalone=True),
]
