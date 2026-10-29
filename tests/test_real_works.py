# -*- coding: utf-8 -*-
"""
Случаи из реальных магистерских ВКР (25 работ, октябрь 2026), на которых проверка ошибалась:
титульный лист со строкой в стиле «Заголовок 1», СОДЕРЖАНИЕ в стиле «Название», другие названия
перечня сокращений и списка источников, «Приложение А» обычным текстом, «Продолжение таблицы»,
столбец сокращений, почти чёрный цвет текста, код обычным шрифтом с комментариями.
"""

import io

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import make_docs as md
from checker import run_normocontrol


def check(doc, tmp_path, profile="vkr"):
    path = tmp_path / "work.docx"
    doc.save(path)
    return run_normocontrol(str(path), profile).issues


def front_page(doc):
    """Титульный лист ПГУПС: название вуза в стиле «Заголовок 1», мелкий кегль, по центру."""
    doc.add_paragraph("ФЕДЕРАЛЬНОЕ АГЕНТСТВО ЖЕЛЕЗНОДОРОЖНОГО ТРАНСПОРТА").alignment = WD_ALIGN_PARAGRAPH.CENTER
    h = doc.add_heading("«Петербургский государственный университет путей сообщения Императора Александра I»", 1)
    h.paragraph_format.page_break_before = False
    for line in ("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА", "Санкт-Петербург", "2026"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line).font.size = Pt(11)


def with_front(build_rest):
    doc = md.new_document()
    front_page(doc)
    build_rest(doc)
    return doc


def body_parts(doc):
    md.structural(doc, "ВВЕДЕНИЕ")
    md.body(doc)
    doc.add_heading("1 Анализ", level=1)
    md.body(doc)
    md.structural(doc, "ЗАКЛЮЧЕНИЕ")
    md.body(doc)
    md.structural(doc, "СПИСОК ИСПОЛЬЗУЕМЫХ ИСТОЧНИКОВ")
    md.body(doc, "1. Иванов, И. И. Методы анализа данных / И. И. Иванов. – Москва : Наука, 2020. – 240 с.")


def test_title_page_with_heading_style_is_front_matter(tmp_path):
    doc = with_front(lambda d: (md.toc(d), body_parts(d)))
    front = [i for i in check(doc, tmp_path) if "ФЕДЕРАЛЬНОЕ" in i.location or "Петербургский" in i.location
             or "Санкт-Петербург" in i.location]
    assert not front, [i.message for i in front]


def test_contents_in_title_style_is_structural(tmp_path):
    def rest(d):
        d.add_paragraph("СОДЕРЖАНИЕ", style="Title")
        for line in ("Введение\t3", "1 Анализ\t4", "Заключение\t5"):
            d.add_paragraph(line).paragraph_format.first_line_indent = Cm(0)
        body_parts(d)
    msgs = [(i.location, i.message) for i in check(with_front(rest), tmp_path)]
    assert not any("ФЕДЕРАЛЬНОЕ" in loc for loc, _ in msgs), msgs[:5]
    assert not any("СОДЕРЖАНИЕ" in m and "не найден" in m for _, m in msgs)


def test_other_names_of_abbreviations_and_sources(tmp_path):
    def rest(d):
        md.toc(d)
        md.structural(d, "СПИСОК СОКРАЩЕНИЙ И УСЛОВНЫХ ОБОЗНАЧЕНИЙ")
        for line in ("НДС — налог на добавленную стоимость", "ИД — индекс доходности"):
            d.add_paragraph(line).paragraph_format.first_line_indent = Cm(0)   # столбцом, без отступа
        body_parts(d)
    msgs = [(i.location, i.category, i.message) for i in check(with_front(rest), tmp_path)]
    assert not any("ФЕДЕРАЛЬНОЕ" in loc for loc, _, _ in msgs)
    assert not any(cat == "Красная строка" and ("НДС" in loc or "ИД —" in loc) for loc, cat, _ in msgs)
    assert not any("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ" in m and "не найден" in m for _, _, m in msgs)


def test_plain_appendix_title_is_recognized(tmp_path):
    doc = md.good_document()
    md.body(doc, "Приложение А").paragraph_format.page_break_before = True
    md.body(doc, "Листинг программы")
    msgs = [(i.location, i.category, i.message) for i in check(doc, tmp_path)]
    assert not any(cat == "Красная строка" and "Приложение А" in loc for loc, cat, _ in msgs)
    assert any("Приложение А" in loc and cat == "Оформление заголовка" for loc, cat, _ in msgs)   # не по центру


def test_table_continuation_label(tmp_path):
    doc = md.good_document()
    p = doc.add_paragraph("Продолжение таблицы 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    bad = doc.add_paragraph("Продолжение таблицы 2")      # с красной строкой — ошибка
    msgs = [(i.location, i.category, i.message) for i in check(doc, tmp_path)]
    assert not any("Продолжение таблицы 1" in loc for loc, _, _ in msgs), msgs
    assert [m for loc, _, m in msgs if "Продолжение таблицы 2" in loc] == \
        ["надпись «Продолжение таблицы» пишется слева, без абзацного отступа (п. 6.6.6)"]


def test_near_black_text_is_black(tmp_path):
    doc = md.good_document()
    for color in ((0x0F, 0x11, 0x15), (0x1A, 0x1A, 0x1A)):
        md.body(doc).runs[0].font.color.rgb = RGBColor(*color)
    md.body(doc).runs[0].font.color.rgb = RGBColor(0x0E, 0x28, 0x41)   # тёмно-синий — уже не чёрный
    colors = [i.message for i in check(doc, tmp_path) if i.category == "Цвет шрифта"]
    assert colors == ["цвет текста #0E2841, ожидается чёрный"]


def test_code_in_body_font_with_comments(tmp_path):
    doc = md.good_document()
    for line in ("def semantic_chunking(text, max_chunk_size=512):",
                 '    """Разбивает текст на смысловые куски"""',
                 "    # Сначала делим по пустым строкам",
                 "    chunks = []",
                 '    current = ""',
                 "    for paragraph in text.split('\\n\\n'):",
                 "        current += paragraph",
                 "    return chunks"):
        p = md.body(doc, line)           # как в работе: Times New Roman, без отступа, по левому краю
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(0)
    body_like = [i for i in check(doc, tmp_path)
                 if i.category in ("Красная строка", "Выравнивание") and ("chunk" in i.location or "Сначала" in i.location
                                                                           or "current" in i.location)]
    assert not body_like, [(i.location, i.message) for i in body_like]


def test_formula_typed_as_text_and_appendix_continuation(tmp_path):
    doc = md.good_document()
    for line in ("E = a + b + c,\t(6.2)", ".                                       (13)", "Продолжение приложения Ж"):
        p = md.body(doc, line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    msgs = [(i.location, i.category) for i in check(doc, tmp_path)
            if i.category in ("Красная строка", "Выравнивание")]
    assert not msgs, msgs
