# -*- coding: utf-8 -*-
"""Копия работы с отметками: примечания на местах ошибок, текст не меняется."""

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Cm, RGBColor

import app as server
import make_docs as md
from checker import run_normocontrol
from checker.annotate import annotate_docx


def broken():
    doc = md.good_document()
    doc.sections[0].left_margin = Cm(2)                    # к разделу документа
    p = doc.add_paragraph()
    p.add_run("Красный текст в основной части работы.").font.color.rgb = RGBColor(255, 0, 0)
    doc.add_paragraph("Текст с  двойным пробелом и красной строкой неправильной длины.") \
        .paragraph_format.first_line_indent = Cm(0.5)        # два замечания в одном абзаце
    doc.add_paragraph("")
    doc.add_paragraph("")                                    # пустые строки — пустой абзац
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 1).paragraphs[0].add_run("Arial в таблице").font.name = "Arial"
    return doc


def annotate(doc, tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "marked.docx"
    doc.save(src)
    report = run_normocontrol(str(src), "vkr")
    n = annotate_docx(str(src), str(dst), report)
    return report, n, Document(str(src)), Document(str(dst))


def test_every_issue_lands_in_a_comment(tmp_path):
    report, n, src, dst = annotate(broken(), tmp_path)
    comments = list(dst.comments)
    assert n == len(comments) > 0
    text = "\n".join(c.text for c in comments)
    for issue in report.issues:
        assert issue.message in text, issue


def test_two_issues_in_one_paragraph_share_one_comment(tmp_path):
    _, _, _, dst = annotate(broken(), tmp_path)
    both = [c for c in dst.comments if "двойной" in c.text]
    assert len(both) == 1 and "Красная строка" in both[0].text


def test_text_unchanged_and_issues_highlighted(tmp_path):
    _, _, src, dst = annotate(broken(), tmp_path)
    assert [p.text for p in src.paragraphs] == [p.text for p in dst.paragraphs]
    red = next(p for p in dst.paragraphs if p.text.startswith("Красный текст"))
    assert all(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in red.runs if r.text)
    ok = next(p for p in dst.paragraphs if p.text.startswith("Современные методы"))
    assert all(r.font.highlight_color is None for r in ok.runs)


def test_table_and_section_issues_anchored(tmp_path):
    _, _, _, dst = annotate(broken(), tmp_path)
    texts = [c.text for c in dst.comments]
    assert any("Шрифт (таблица)" in t for t in texts)
    assert any(t.startswith("Раздел документа №1:") and "левое поле" in t for t in texts)


def test_marked_file_does_not_add_issues(tmp_path):
    report, _, _, _ = annotate(broken(), tmp_path)
    again = run_normocontrol(str(tmp_path / "marked.docx"), "vkr")
    assert len(again.issues) == len(report.issues)


def test_multi_section_document(tmp_path):
    doc = md.good_document()
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.footer.is_linked_to_previous = False                 # раздел 2 без номеров страниц
    md.body(doc)
    _, _, _, dst = annotate(doc, tmp_path)
    assert any(c.text.startswith("Раздел документа №2:") for c in dst.comments)


def test_api_annotated_download():
    buf = io.BytesIO()
    broken().save(buf)
    c = server.app.test_client()
    rep = c.post("/api/check", data={"profile": "vkr", "file": (io.BytesIO(buf.getvalue()), "работа.docx")},
                 content_type="multipart/form-data").get_json()
    r = c.get(f"/api/annotated/{rep['upload_id']}?profile=vkr")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert "attachment" in r.headers["Content-Disposition"]
    assert len(list(Document(io.BytesIO(r.data)).comments)) > 0
    assert c.get("/api/annotated/" + "0" * 32 + "?profile=vkr").status_code == 404
