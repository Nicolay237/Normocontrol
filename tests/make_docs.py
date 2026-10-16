# -*- coding: utf-8 -*-
"""Генератор тестовых .docx: эталонная работа и её намеренно испорченные варианты."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

LOREM = (
    "Современные методы анализа данных позволяют существенно повысить качество "
    "принимаемых решений при проектировании информационных систем, поэтому в данной "
    "работе рассматривается подход, основанный на последовательной проверке гипотез. "
) * 3


def _page_field(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def set_style_font(style, name="Times New Roman", size=14):
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), name)


def new_document(left=3.0, right=1.5, top=2.0, bottom=2.0, page_numbers=True,
                 body_size=14, body_font="Times New Roman", spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                 indent=1.25):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin, s.right_margin = Cm(left), Cm(right)
    s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)

    normal = doc.styles["Normal"]
    set_style_font(normal, body_font, body_size)
    pf = normal.paragraph_format
    pf.line_spacing = spacing
    pf.alignment = align
    pf.first_line_indent = Cm(indent)
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    for name in ("Heading 1", "Heading 2"):
        st = doc.styles[name]
        set_style_font(st, "Times New Roman", 14)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.first_line_indent = Cm(1.25)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.line_spacing = 1.5
    doc.styles["Heading 1"].paragraph_format.page_break_before = True

    if page_numbers:
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _page_field(fp)
    return doc


def structural(doc, text, page_break=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    if page_break:
        p.paragraph_format.page_break_before = True
    return p


def body(doc, text=LOREM):
    return doc.add_paragraph(text)


def good_document():
    doc = new_document()
    structural(doc, "ВВЕДЕНИЕ", page_break=False)
    for _ in range(3):
        body(doc)
    doc.add_heading("1 Теоретические основы", level=1)
    body(doc)
    doc.add_heading("1.1 Анализ предметной области", level=2)
    body(doc)
    body(doc)
    # рисунок и подпись
    img = doc.add_paragraph()
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img.paragraph_format.first_line_indent = Cm(0)
    img.add_run().add_picture(_png(), width=Cm(4))
    cap = doc.add_paragraph("Рисунок 1 – Схема работы системы")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing = 1.0
    # таблица и название
    t_cap = doc.add_paragraph("Таблица 1 – Результаты эксперимента")
    t_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    t_cap.paragraph_format.first_line_indent = Cm(0)
    t_cap.paragraph_format.line_spacing = 1.0
    table = doc.add_table(rows=2, cols=2)
    for row in table.rows:
        for c in row.cells:
            c.paragraphs[0].add_run("1,5")
    body(doc)
    # список через тире
    for item in ("- первый пункт перечисления;", "- второй пункт перечисления."):
        lp = doc.add_paragraph(item)
        lp.paragraph_format.first_line_indent = Cm(1.25)
    structural(doc, "ЗАКЛЮЧЕНИЕ")
    for _ in range(3):
        body(doc)
    structural(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    body(doc, "1. Иванов И. И. Методы анализа данных. – М.: Наука, 2020. – 240 с.")
    return doc


def _png():
    import io, base64
    data = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    return io.BytesIO(data)
