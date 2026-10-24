# -*- coding: utf-8 -*-
"""Подписи без таблиц и рисунков, подписи не на своём месте, обязательное СОДЕРЖАНИЕ (найдено тестировщиками)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

import make_docs as md
from checker import run_normocontrol
from checker.fixers import fix_docx
from checker.norms import PROFILES


def caption(doc, text, center=False):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def table(doc):
    t = doc.add_table(rows=2, cols=2)
    for row in t.rows:
        for c in row.cells:
            c.paragraphs[0].add_run("1")


def picture(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(md._png(), width=Cm(4))


def work(build):
    doc = md.new_document()
    md.toc(doc)
    md.structural(doc, "ВВЕДЕНИЕ")
    md.body(doc)
    build(doc)
    md.body(doc)
    md.structural(doc, "ЗАКЛЮЧЕНИЕ")
    md.body(doc)
    md.structural(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    md.body(doc, "1. Иванов, И. И. Методы анализа данных / И. И. Иванов. – Москва : Наука, 2020. – 240 с.")
    return doc


def messages(doc, tmp_path):
    path = tmp_path / "w.docx"
    doc.save(path)
    return [i.message for i in run_normocontrol(str(path), "vkr").issues]


def test_table_caption_without_table(tmp_path):
    msgs = messages(work(lambda d: caption(d, "Таблица 1 – Результаты эксперимента")), tmp_path)
    assert any("нет самой таблицы" in m for m in msgs)


def test_figure_caption_without_picture(tmp_path):
    msgs = messages(work(lambda d: caption(d, "Рисунок 1 – Схема системы", center=True)), tmp_path)
    assert any("над подписью нет рисунка" in m for m in msgs)


def test_sentence_about_table_is_not_a_caption(tmp_path):
    msgs = messages(work(lambda d: md.body(d, "Таблица 1 показывает, что метод работает быстрее аналогов.")), tmp_path)
    assert not any("таблиц" in m for m in msgs)


def misplaced(doc):
    table(doc)
    caption(doc, "Таблица 1 – Результаты эксперимента")
    caption(doc, "Рисунок 1 – Схема системы", center=True)
    picture(doc)


def test_captions_on_the_wrong_side(tmp_path):
    msgs = messages(work(misplaced), tmp_path)
    assert any("стоит под таблицей" in m for m in msgs)
    assert any("стоит над рисунком" in m for m in msgs)
    assert not any("нет самой таблицы" in m or "нет рисунка" in m for m in msgs)


def test_autofix_moves_captions(tmp_path):
    src, dst = tmp_path / "src.docx", tmp_path / "dst.docx"
    work(misplaced).save(src)
    result = fix_docx(str(src), str(dst), PROFILES["vkr"])
    assert not [i for i in result.after.issues if "стоит под" in i.message or "стоит над" in i.message]
    body = Document(dst).element.body
    tags = [(el.tag.split("}")[1], "".join(el.itertext())[:20]) for el in body if el.tag.split("}")[1] in ("p", "tbl")]
    t = next(k for k, (tag, _) in enumerate(tags) if tag == "tbl")
    assert tags[t - 1][1].startswith("Таблица 1")              # название — над таблицей
    pic = next(k for k, (tag, text) in enumerate(tags) if tag == "p" and text.startswith("Рисунок 1"))
    assert tags[pic - 1][1] == ""                               # подпись — сразу под абзацем с рисунком


def test_toc_is_required(tmp_path):
    doc = md.new_document()
    md.structural(doc, "ВВЕДЕНИЕ", page_break=False)
    for _ in range(6):
        md.body(doc)
    md.structural(doc, "ЗАКЛЮЧЕНИЕ")
    md.body(doc)
    assert any("«СОДЕРЖАНИЕ»" in m for m in messages(doc, tmp_path))
