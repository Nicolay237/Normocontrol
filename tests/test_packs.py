# -*- coding: utf-8 -*-
"""Дополнительные проверки вуза: титульный лист и задание ПГУПС, доступ через администратора."""

import io
from datetime import date

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt

import app as server
import make_docs as md
from checker import run_normocontrol

from test_auth import ADMIN, admin_client, client, login, register

HEADER = [
    ("ФЕДЕРАЛЬНОЕ АГЕНТСТВО ЖЕЛЕЗНОДОРОЖНОГО ТРАНСПОРТА", 10, False),
    ("Федеральное государственное бюджетное образовательное учреждение высшего образования", 10, False),
    ("«Петербургский государственный университет путей сообщения", 12, True),
    ("Императора Александра I»", 12, True),
    ("(ФГБОУ ВО ПГУПС)", 12, True),
]
YEAR = date.today().year


def line(doc, text="", size=12, bold=False, center=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def pgups_document(student="Корнеев Даниил Борисович", head="д.т.н., профессор Ермаков С. Г.",
                   supervisor="д.ф.-м.н., профессор Бригаднов И. А.", year=str(YEAR), topic="«Разработка системы»",
                   task=True, speciality="09.03.01 «Информатика и вычислительная техника»", theme="Разработка системы",
                   order="от «12» марта 2026 г. № 345", plan=True, abstract="АННОТАЦИЯ"):
    doc = md.new_document()
    for text, size, bold in HEADER:
        line(doc, text, size, bold, center=True)
    line(doc)
    line(doc, "Факультет «Автоматизация и интеллектуальные технологии»")
    line(doc, "Кафедра «Информационные и вычислительные системы»")
    line(doc)
    line(doc, "ПОЯСНИТЕЛЬНАЯ ЗАПИСКА", 16, True, center=True)
    line(doc, "К ВЫПУСКНОЙ КВАЛИФИКАЦИОННОЙ РАБОТЕ", 16, True, center=True)
    line(doc)
    line(doc, student, center=True)
    line(doc, "Фамилия, имя, отчество обучающегося", 8)
    line(doc, f"на тему {topic}")
    t = doc.add_table(rows=4, cols=3)
    for row, (label, who) in zip(t.rows, [("Обучающийся", "Корнеев Д. Б."), ("Заведующий кафедрой", head),
                                           ("Руководитель ВКР", supervisor),
                                           ("Нормоконтролер", "к.т.н., доцент Петров В. Е.")]):
        row.cells[0].text, row.cells[1].text, row.cells[2].text = label, "__________ подпись, дата", who
    line(doc)
    line(doc, f"Санкт-Петербург {year}", center=True)

    if task:
        page_break(doc)
        for text, size, bold in HEADER:
            line(doc, text, size, bold, center=True)
        line(doc, "Факультет «Автоматизация»     Кафедра «Информационные системы»")
        line(doc, f"Специальность (направление) {speciality}")
        line(doc, "Специализация (профиль) «Программное обеспечение»")
        line(doc, "УТВЕРЖДАЮ")
        line(doc, "Зав. кафедрой")
        line(doc, "«15» марта 2026 г.")
        line(doc, "ЗАДАНИЕ", 16, True, center=True)
        line(doc, "на выпускную квалификационную работу обучающегося", 11, True, center=True)
        line(doc, "Корнееву Даниилу Борисовичу", center=True)
        line(doc, "фамилия, имя, отчество", 9, center=True)
        line(doc, f"1. Тема ВКР {theme}")
        line(doc, f"утверждена распоряжением первого проректора – проректора по учебной работе {order}")
        line(doc, "2. Срок сдачи обучающимся законченной ВКР 1 июня 2026 г.")
        line(doc, "3. Исходные данные к ВКР ____________")
        line(doc, "Материалы преддипломной практики")
        line(doc, "4. Содержание пояснительной записки ВКР (перечень подлежащих разработке вопросов)")
        line(doc, "Анализ, проектирование, реализация")
        line(doc, "5. Перечень графического материала (с точным указанием обязательных чертежей / слайдов)")
        line(doc, "Презентация, 12 слайдов")
        line(doc, "6. Дата выдачи задания 15 марта 2026 г.")
        line(doc, "Руководитель ВКР ________ ________")
        line(doc, "Задание принял к исполнению ________ ________")
        page_break(doc)
        line(doc, "КАЛЕНДАРНЫЙ ПЛАН", bold=True)
        cal = doc.add_table(rows=3, cols=4)
        for c, text in zip(cal.rows[0].cells, ("№ п/п", "Наименование этапов ВКР", "Сроки", "Примечание")):
            c.text = text
        if plan:
            for c, text in zip(cal.rows[1].cells, ("1", "Обзор литературы", "до 01.04.2026", "")):
                c.text = text

    md.structural(doc, abstract)
    md.body(doc)
    md.structural(doc, "СОДЕРЖАНИЕ")
    md.body(doc, "Введение 4")
    md.structural(doc, "ВВЕДЕНИЕ")
    md.body(doc)
    md.structural(doc, "ЗАКЛЮЧЕНИЕ")
    md.body(doc)
    md.structural(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    md.body(doc, "1. Иванов, И. И. Методы анализа данных / И. И. Иванов. – Москва : Наука, 2020. – 240 с.")
    return doc


def check(doc, tmp_path, packs=(("pgups", None),)):
    path = tmp_path / "w.docx"
    doc.save(path)
    return run_normocontrol(str(path), "vkr", packs=packs)


def pgups_messages(report):
    return [i.message for i in report.issues if "ПГУПС" in i.category]


def test_correct_front_matter_passes(tmp_path):
    report = check(pgups_document(), tmp_path)
    assert pgups_messages(report) == []
    assert any("ПГУПС" in n for n in report.notes)


def test_title_page_is_not_checked_as_body_text(tmp_path):
    report = check(pgups_document(), tmp_path, packs=())
    title_text = ("ФЕДЕРАЛЬНОЕ", "Факультет", "ПОЯСНИТЕЛЬНАЯ", "Корнеев", "Санкт-Петербург", "ЗАДАНИЕ", "Тема ВКР")
    assert not [i for i in report.issues if any(t in i.location for t in title_text)]
    assert not [i for i in report.issues if i.category == "Шрифт (таблица)"]


def test_title_page_errors(tmp_path):
    msgs = pgups_messages(check(pgups_document(student="Корнеев Д. Б.", head="Ермаков С. Г.",
                                               supervisor="профессор", year="20__", topic="______"), tmp_path))
    joined = "\n".join(msgs)
    assert "ФИО обучающегося пишется полностью" in joined
    assert "«Заведующий кафедрой»: не указаны учёная степень или звание" in joined
    assert "«Руководитель ВКР»: не указаны фамилия и инициалы" in joined
    assert "не указана тема ВКР" in joined
    assert "не указан год" in joined


def test_task_errors(tmp_path):
    msgs = "\n".join(pgups_messages(check(pgups_document(
        speciality="«Информатика»", theme="____", order="от «__» _____ 20__ г. № ___", plan=False), tmp_path)))
    assert "нет шифра по ФГОС" in msgs
    assert "не заполнен пункт «1. Тема ВКР»" in msgs
    assert "дата и номер распоряжения" in msgs
    assert "календарный план не заполнен" in msgs


def test_missing_task_and_wrong_names(tmp_path):
    msgs = "\n".join(pgups_messages(check(pgups_document(task=False, abstract="РЕФЕРАТ"), tmp_path)))
    assert "не найдено задание на ВКР" in msgs
    assert "вместо реферата — «АННОТАЦИЯ»" in msgs


def test_pack_settings_override_defaults(tmp_path):
    report = check(pgups_document(task=False), tmp_path, packs=(("pgups", '{"require_task": false}'),))
    assert pgups_messages(report) == []


def test_pack_only_for_vkr(tmp_path):
    path = tmp_path / "w.docx"
    pgups_document(task=False).save(path)
    report = run_normocontrol(str(path), "coursework", packs=(("pgups", None),))
    assert pgups_messages(report) == []


# ── доступ: только по назначению администратора ──

def docx_upload(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return {"profile": "vkr", "file": (io.BytesIO(buf.getvalue()), "w.docx")}


def test_pack_applies_only_after_admin_assigns_it():
    c = client()
    register(c, "pgups@mail.ru", ip="10.1.1.1")
    doc = pgups_document(student="Корнеев Д. Б.")
    r = c.post("/api/check", data=docx_upload(doc), content_type="multipart/form-data").get_json()
    assert r["extra_checks"] == [] and not any("ПГУПС" in cat["name"] for cat in r["categories"])

    admin = admin_client()
    me = c.get("/api/me").get_json()["user"]
    assert admin.patch(f"/api/admin/users/{me['id']}", json={"packs": ["nope"]}).status_code == 400
    assert admin.patch(f"/api/admin/users/{me['id']}", json={"packs": ["pgups"]}).status_code == 200
    assert [p["key"] for p in c.get("/api/me").get_json()["user"]["packs"]] == ["pgups"]

    r = c.post("/api/check", data=docx_upload(doc), content_type="multipart/form-data").get_json()
    assert r["extra_checks"] and any(cat["name"] == "Титульный лист (ПГУПС)" for cat in r["categories"])
    row = next(u for u in admin.get("/api/admin/users").get_json()["users"] if u["id"] == me["id"])
    assert row["packs"] == ["pgups"]

    admin.patch(f"/api/admin/users/{me['id']}", json={"packs": []})
    r = c.post("/api/check", data=docx_upload(doc), content_type="multipart/form-data").get_json()
    assert r["extra_checks"] == []


def test_pack_settings_admin_api():
    admin = admin_client()
    packs = admin.get("/api/admin/packs").get_json()["packs"]
    pg = next(p for p in packs if p["key"] == "pgups")
    assert pg["settings"]["city"] == "Санкт-Петербург" and pg["customized"] is False

    assert admin.put("/api/admin/packs/pgups", json={"settings": {"unknown": 1}}).status_code == 400
    assert admin.put("/api/admin/packs/pgups", json={"settings": {"require_task": "no"}}).status_code == 400
    r = admin.put("/api/admin/packs/pgups", json={"settings": {**pg["settings"], "require_task": False}})
    assert r.status_code == 200 and r.get_json()["pack"]["customized"] is True
    assert r.get_json()["pack"]["settings"]["require_task"] is False
    r = admin.put("/api/admin/packs/pgups", json={"settings": None})      # сброс к значениям по умолчанию
    assert r.get_json()["pack"]["customized"] is False

    c = client()
    register(c, "nosy@mail.ru", ip="10.1.1.2")
    assert c.get("/api/admin/packs").status_code == 403
